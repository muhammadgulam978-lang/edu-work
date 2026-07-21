from datetime import timedelta
from io import BytesIO
from urllib.parse import quote

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.mail import EmailMessage
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST

from admin_panel.models import Class, Section
from student_profile.models import Student

from .agent import run_admin_agent
from .models import AdminAIConversation, AdminAIMessage, AdminAIReport, AdminAIReportShareLog
from .services import (
    PERIOD_LABELS,
    build_advanced_analytics,
    get_student_intelligence,
    make_json_safe,
    parse_filter_request,
)


REPORT_SECTIONS = [
    ("kpis", "KPI Cards"),
    ("summary", "Summary"),
    ("charts", "Charts"),
    ("tables", "Tables"),
    ("student_risk", "Student Risk"),
    ("ai_tutor", "AI Tutor Usage"),
]


def _admin_context():
    return {
        "period_options": list(PERIOD_LABELS.items()),
        "classes": Class.objects.all().order_by("class_name"),
        "sections": Section.objects.select_related("class_fk").order_by("class_fk__class_name", "section_name"),
    }


@login_required
@permission_required("auth.view_group", raise_exception=True)
def admin_ai_copilot(request):
    conversation = (
        AdminAIConversation.objects.filter(created_by=request.user)
        .prefetch_related("messages")
        .first()
    )
    if not conversation:
        conversation = AdminAIConversation.objects.create(created_by=request.user)
    return render(request, "admin_ai/copilot.html", {"conversation": conversation})


@login_required
@permission_required("auth.view_group", raise_exception=True)
@require_POST
def admin_ai_copilot_message(request):
    question = request.POST.get("message", "").strip()
    if not question:
        return JsonResponse({"error": "Enter a question for EduPilot AI Copilot."}, status=400)
    conversation_id = request.POST.get("conversation_id")
    conversation = AdminAIConversation.objects.filter(id=conversation_id, created_by=request.user).first()
    if not conversation:
        conversation = AdminAIConversation.objects.create(created_by=request.user)

    AdminAIMessage.objects.create(conversation=conversation, sender="admin", content=question)
    result = run_admin_agent(query=question, user=request.user, conversation=conversation)
    message = AdminAIMessage.objects.create(
        conversation=conversation,
        sender="assistant",
        content=result["answer"],
        intent=result["intent"],
        tool_name=", ".join([tool.get("name", "") for tool in result.get("tools", [])])[:120],
        payload=make_json_safe(result),
    )
    return JsonResponse({
        "reply": result["answer"],
        "intent": result["intent"],
        "tool": message.tool_name,
        "plan": result["plan"],
        "tools": result["tools"],
        "evidence": result["evidence"],
        "actions": result["actions"],
        "mode": result["mode"],
        "model": result["model"],
        "llm_error": result["llm_error"],
        "message_id": message.id,
    })


@login_required
@permission_required("auth.view_group", raise_exception=True)
def ai_analytics_advanced(request):
    filters = parse_filter_request(request)
    data = build_advanced_analytics(filters)
    context = _admin_context()
    context.update({"analytics": data})
    return render(request, "admin_ai/advanced_analytics.html", context)


@login_required
@permission_required("auth.view_group", raise_exception=True)
def ai_analytics_advanced_data(request):
    return JsonResponse(build_advanced_analytics(parse_filter_request(request)))


@login_required
@permission_required("auth.view_group", raise_exception=True)
def ai_analytics_reports(request):
    context = _admin_context()
    context.update({
        "report_types": AdminAIReport.REPORT_TYPES,
        "report_sections": REPORT_SECTIONS,
        "recent_reports": AdminAIReport.objects.select_related("generated_by")[:30],
    })
    return render(request, "admin_ai/reports.html", context)


@login_required
@permission_required("auth.view_group", raise_exception=True)
@require_POST
def ai_analytics_report_generate(request):
    report_type = request.POST.get("report_type") or "full"
    title = request.POST.get("title") or f"EduPilot AI Analytics Report - {timezone.localdate()}"
    filters = parse_filter_request(request)
    sections = request.POST.getlist("sections") or [key for key, _label in REPORT_SECTIONS]
    snapshot = build_advanced_analytics(filters)
    snapshot["report"] = {
        "type": report_type,
        "sections": sections,
        "generated_at": timezone.now().isoformat(),
    }
    report = AdminAIReport.objects.create(
        title=title,
        report_type=report_type,
        generated_by=request.user,
        snapshot_data=make_json_safe(snapshot),
        expires_at=timezone.now() + timedelta(days=30),
    )
    messages.success(request, "AI Analytics report generated.")
    return redirect("ai_analytics_report_detail", report_uuid=report.uuid)


@login_required
@permission_required("auth.view_group", raise_exception=True)
@require_POST
def ai_analytics_report_auto_generate(request):
    report_type = request.POST.get("report_type") or "full"
    filters = parse_filter_request(request)
    sections = request.POST.getlist("sections") or [key for key, _label in REPORT_SECTIONS]
    period_label = PERIOD_LABELS.get(filters.get("period"), filters.get("period") or "Selected period")
    title = request.POST.get("title") or f"EduPilot AI Analytics {period_label} Report - {timezone.localdate()}"
    source = request.POST.get("source") or "auto_report_button"

    snapshot = build_advanced_analytics(filters)
    snapshot["report"] = {
        "type": report_type,
        "sections": sections,
        "generated_at": timezone.now().isoformat(),
        "source": source,
        "auto_generated": True,
    }
    report = AdminAIReport.objects.create(
        title=title,
        report_type=report_type,
        generated_by=request.user,
        snapshot_data=make_json_safe(snapshot),
        expires_at=timezone.now() + timedelta(days=30),
    )
    detail_url = reverse("ai_analytics_report_detail", kwargs={"report_uuid": report.uuid})
    return JsonResponse({
        "success": True,
        "title": report.title,
        "message": "Your report is ready. You can download the PDF, open the preview, or share it.",
        "report_url": detail_url,
        "pdf_url": reverse("ai_analytics_report_pdf", kwargs={"report_uuid": report.uuid}),
        "excel_url": reverse("ai_analytics_report_excel", kwargs={"report_uuid": report.uuid}),
        "share_url": f"{detail_url}#share-report",
    })


@login_required
@permission_required("auth.view_group", raise_exception=True)
def ai_analytics_report_detail(request, report_uuid):
    report = get_object_or_404(AdminAIReport, uuid=report_uuid)
    report_url = request.build_absolute_uri(reverse("ai_analytics_report_detail", kwargs={"report_uuid": report.uuid}))
    social_text = f"EduPilot report: {report.title}"
    encoded_text = quote(social_text)
    encoded_url = quote(report_url)
    return render(request, "admin_ai/report_detail.html", {
        "report": report,
        "snapshot": report.snapshot_data,
        "report_url": report_url,
        "whatsapp_url": f"https://wa.me/?text={encoded_text}%20{encoded_url}",
        "facebook_url": f"https://www.facebook.com/sharer/sharer.php?u={encoded_url}",
        "x_url": f"https://twitter.com/intent/tweet?text={encoded_text}&url={encoded_url}",
    })


def _report_pdf_bytes(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=32, rightMargin=32, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("EduPilot AI Analytics Report", styles["Title"]),
        Paragraph(report.title, styles["Heading2"]),
        Paragraph(f"Generated: {report.created_at:%B %d, %Y, %I:%M %p}", styles["Normal"]),
        Spacer(1, 14),
    ]
    snapshot = report.snapshot_data or {}
    summary = snapshot.get("summary", {})
    kpis = snapshot.get("kpis", [])
    alerts_rows = [["Metric", "Value"]]
    for key, value in summary.items():
        alerts_rows.append([key.replace("_", " ").title(), str(value)])
    if len(alerts_rows) > 1:
        story.append(Paragraph("Summary", styles["Heading3"]))
        story.append(_pdf_table(alerts_rows))
        story.append(Spacer(1, 12))
    if kpis:
        story.append(Paragraph("KPIs", styles["Heading3"]))
        story.append(_pdf_table([["Title", "Value", "Label"]] + [[k.get("title"), k.get("value"), k.get("label")] for k in kpis]))
        story.append(Spacer(1, 12))
    tables = snapshot.get("tables", {})
    for title, rows in tables.items():
        if rows:
            keys = list(rows[0].keys())
            story.append(Paragraph(title.replace("_", " ").title(), styles["Heading3"]))
            story.append(_pdf_table([[key.replace("_", " ").title() for key in keys]] + [[str(row.get(key, "")) for key in keys] for row in rows[:20]]))
            story.append(Spacer(1, 12))
    doc.build(story)
    return buffer.getvalue()


def _pdf_table(rows):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#88D0C5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e9ecef")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    return table


@login_required
@permission_required("auth.view_group", raise_exception=True)
def ai_analytics_report_pdf(request, report_uuid):
    report = get_object_or_404(AdminAIReport, uuid=report_uuid)
    pdf = _report_pdf_bytes(report)
    response = HttpResponse(pdf, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="EduPilot_AI_Analytics_Report_{report.created_at:%Y-%m-%d}.pdf"'
    AdminAIReportShareLog.objects.create(report=report, shared_by=request.user, channel="download", status="downloaded")
    return response


@login_required
@permission_required("auth.view_group", raise_exception=True)
def ai_analytics_report_excel(request, report_uuid):
    import openpyxl

    report = get_object_or_404(AdminAIReport, uuid=report_uuid)
    wb = openpyxl.Workbook()
    snapshot = report.snapshot_data or {}
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Report", report.title])
    ws.append(["Generated", report.created_at.strftime("%Y-%m-%d %H:%M")])
    ws.append([])
    ws.append(["Metric", "Value"])
    for key, value in snapshot.get("summary", {}).items():
        ws.append([key.replace("_", " ").title(), value])

    _add_sheet(wb, "KPIs", [["Title", "Value", "Label"]] + [[k.get("title"), k.get("value"), k.get("label")] for k in snapshot.get("kpis", [])])
    _add_chart_sheet(wb, "Attendance", snapshot.get("charts", {}).get("attendance", {}))
    _add_chart_sheet(wb, "Fees", snapshot.get("charts", {}).get("fees", {}))
    _add_chart_sheet(wb, "Admissions", snapshot.get("charts", {}).get("admissions", {}))
    _add_chart_sheet(wb, "Fixtures", snapshot.get("charts", {}).get("fixture_mix", {}))
    _add_chart_sheet(wb, "Teacher Workload", snapshot.get("charts", {}).get("teacher_workload", {}))
    for title, rows in snapshot.get("tables", {}).items():
        if rows:
            keys = list(rows[0].keys())
            _add_sheet(wb, title[:31], [[key.replace("_", " ").title() for key in keys]] + [[row.get(key, "") for key in keys] for row in rows])

    buffer = BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="EduPilot_AI_Analytics_Report_{report.created_at:%Y-%m-%d}.xlsx"'
    AdminAIReportShareLog.objects.create(report=report, shared_by=request.user, channel="download", status="downloaded")
    return response


def _add_sheet(wb, title, rows):
    ws = wb.create_sheet(title[:31])
    for row in rows:
        ws.append(row)


def _add_chart_sheet(wb, title, chart):
    labels = chart.get("labels", [])
    series_keys = [key for key in chart.keys() if key != "labels"]
    rows = [["Label"] + [key.replace("_", " ").title() for key in series_keys]]
    for index, label in enumerate(labels):
        rows.append([label] + [chart.get(key, [])[index] if index < len(chart.get(key, [])) else "" for key in series_keys])
    _add_sheet(wb, title, rows)


@login_required
@permission_required("auth.view_group", raise_exception=True)
@require_POST
def ai_analytics_report_share(request, report_uuid):
    report = get_object_or_404(AdminAIReport, uuid=report_uuid)
    raw_emails = request.POST.get("emails", "")
    optional_message = request.POST.get("message", "")
    emails = [email.strip() for email in raw_emails.replace("\n", ",").split(",") if email.strip()]
    report_url = request.build_absolute_uri(reverse("ai_analytics_report_detail", kwargs={"report_uuid": report.uuid}))
    if not emails:
        messages.error(request, "Please add at least one recipient email.")
        return redirect("ai_analytics_report_detail", report_uuid=report.uuid)
    try:
        email = EmailMessage(
            subject=f"EduPilot AI Analytics Report: {report.title}",
            body=f"{optional_message}\n\nOpen internal report link: {report_url}",
            from_email=getattr(settings, "DEFAULT_FROM_EMAIL", None),
            to=emails,
        )
        email.attach(f"EduPilot_AI_Analytics_Report_{report.created_at:%Y-%m-%d}.pdf", _report_pdf_bytes(report), "application/pdf")
        email.send(fail_silently=False)
        report.mark_shared(emails)
        AdminAIReportShareLog.objects.create(report=report, shared_by=request.user, channel="email", recipients=emails, status="sent", message=optional_message)
        messages.success(request, "Report shared by email.")
    except Exception as exc:
        report.share_status = "failed"
        report.share_emails = emails
        report.last_shared_at = timezone.now()
        report.save(update_fields=["share_status", "share_emails", "last_shared_at"])
        AdminAIReportShareLog.objects.create(report=report, shared_by=request.user, channel="email", recipients=emails, status="failed", message=optional_message, error_message=str(exc))
        messages.error(request, f"Email failed, but the report link is still available: {report_url}")
    return redirect("ai_analytics_report_detail", report_uuid=report.uuid)


@login_required
@permission_required("auth.view_group", raise_exception=True)
def student_intelligence(request):
    query = request.GET.get("q", "").strip()
    results = []
    if query:
        results = list(Student.objects.filter(student_id__icontains=query)[:25])
    return render(request, "admin_ai/student_intelligence.html", {"query": query, "results": results})


@login_required
@permission_required("auth.view_group", raise_exception=True)
def student_intelligence_detail(request, student_id):
    snapshot = get_student_intelligence(student_id, user=request.user)
    if not snapshot:
        messages.error(request, "Student profile not found.")
        return redirect("admin_ai_student_intelligence")
    profile_url = request.build_absolute_uri(reverse("admin_ai_student_intelligence_detail", kwargs={"student_id": snapshot["student"]["student_id"]}))
    return render(request, "admin_ai/student_intelligence_detail.html", {"snapshot": snapshot, "profile_url": profile_url})


@login_required
@permission_required("auth.view_group", raise_exception=True)
def student_intelligence_pdf(request, student_id):
    snapshot = get_student_intelligence(student_id, user=request.user)
    if not snapshot:
        return HttpResponse("Student not found", status=404)
    fake_report = type("StudentProfileReport", (), {
        "title": f"Student Profile - {snapshot['student']['name']}",
        "created_at": timezone.now(),
        "snapshot_data": {
            "summary": {
                "student_id": snapshot["student"]["student_id"],
                "name": snapshot["student"]["name"],
                "class": snapshot["student"]["class"],
                "section": snapshot["student"]["section"],
                "attendance_rate": snapshot["attendance"]["rate"],
                "unpaid_vouchers": snapshot["fees"]["unpaid"],
                "average_marks": snapshot["marks"]["average"],
            },
            "tables": {
                "recent_attendance": snapshot["attendance"]["recent"],
                "recent_fees": snapshot["fees"]["recent"],
                "recent_marks": snapshot["marks"]["recent"],
            },
        },
    })()
    response = HttpResponse(_report_pdf_bytes(fake_report), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="EduPilot_Student_Profile_{snapshot["student"]["student_id"]}.pdf"'
    return response
