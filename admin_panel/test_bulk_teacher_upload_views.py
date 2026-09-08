from io import BytesIO
from tempfile import TemporaryDirectory

from django.contrib.auth.models import User, Group
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from teacher_dashboard.models import Teacher
from .bulk_teacher_import import import_teachers_from_worksheet


@override_settings(BULK_TEACHER_CHUNK_ROWS=1)
class BulkTeacherUploadViewTests(TestCase):
    def setUp(self):
        self.media = TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        self.override = override_settings(MEDIA_ROOT=self.media.name)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.admin = User.objects.create_superuser('upload.admin', 'admin@example.com', 'StrongPass@12345')
        self.client.force_login(self.admin)

    def upload(self, duplicate=False, invalid=False):
        wb = Workbook()
        wb.active.append(['Name', 'Email', 'Login_Id', 'Password', 'Date Of Birth', 'Status'])
        wb.active.append(['Sample Teacher', 'teacher@example.com', 'teacher.sample', 'StrongPass@12345',
                          'bad-date' if invalid else '1990-05-06', 'Active'])
        if duplicate:
            wb.active.append(['Second Teacher', 'TEACHER@example.com', 'teacher.second', 'StrongPass@12345', '', 'Inactive'])
        stream = BytesIO()
        wb.save(stream)
        return SimpleUploadedFile('teachers.xlsx', stream.getvalue())

    def progress(self, operation, token, cursor=None):
        data = {'operation': operation, 'upload_token': token}
        if cursor is not None:
            data['cursor'] = cursor
        return self.client.post(reverse('bulk_upload_teachers_progress'), data)

    def validate(self, **kwargs):
        response = self.client.post(reverse('bulk_upload_teachers'),
                                    {'excel_file': self.upload(**kwargs)},
                                    HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 202)
        token = response.json()['token']
        state = self.progress('validate_start', token).json()
        while not state['done']:
            response = self.progress('validate_process', token, state['cursor'])
            self.assertEqual(response.status_code, 200)
            state = response.json()
        return token

    def test_complete_chunked_workflow_and_credentials(self):
        token = self.validate()
        self.assertFalse(Teacher.objects.exists())
        self.assertFalse(Group.objects.filter(name='Teacher').exists())
        screen = self.client.get(reverse('bulk_upload_teachers'))
        self.assertContains(screen, 'Validation Preview')
        self.assertContains(screen, 'data-progress-overlay')
        state = self.progress('start', token).json()
        response = self.progress('process', token, state['cursor'])
        self.assertTrue(response.json()['done'])
        teacher = Teacher.objects.get(email='teacher@example.com')
        self.assertTrue(teacher.user.check_password('StrongPass@12345'))
        self.assertTrue(teacher.user.groups.filter(name='Teacher').exists())
        self.assertContains(self.client.get(reverse('bulk_upload_teachers')), 'Complete Import Report')
        self.assertEqual(self.progress('process', token, state['cursor']).status_code, 409)
        response = self.client.get(reverse('bulk_upload_teachers_credentials'))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(b''.join(response.streaming_content)))
        self.assertEqual(workbook.active.cell(2, 2).value, 'teacher.sample')
        response.close()
        self.assertEqual(self.client.get(reverse('bulk_upload_teachers_activity')).json()['teacher_count'], 1)

    def test_duplicate_across_chunks_blocks_all_confirmation_paths(self):
        token = self.validate(duplicate=True)
        self.assertEqual(self.client.session['bulk_teacher_pending']['invalid_rows'], 1)
        self.assertEqual(self.progress('start', token).status_code, 400)
        response = self.client.post(reverse('bulk_upload_teachers'),
                                    {'action': 'confirm', 'upload_token': token}, follow=True)
        self.assertContains(response, 'Only a fully validated file')
        self.assertFalse(Teacher.objects.exists())

    def test_invalid_date_preview_has_no_side_effects(self):
        self.validate(invalid=True)
        self.assertFalse(User.objects.filter(username='teacher.sample').exists())
        self.assertEqual(self.client.session['bulk_teacher_pending']['invalid_rows'], 1)

    def test_stale_cursor_and_unvalidated_import_are_rejected(self):
        response = self.client.post(reverse('bulk_upload_teachers'), {'excel_file': self.upload()},
                                    HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        token = response.json()['token']
        self.assertEqual(self.progress('start', token).status_code, 400)
        self.progress('validate_start', token)
        self.assertEqual(self.progress('validate_process', token, 99).status_code, 409)
        self.assertFalse(Teacher.objects.exists())

    def test_non_javascript_preview_confirm_and_cancel(self):
        response = self.client.post(reverse('bulk_upload_teachers'), {'excel_file': self.upload()})
        self.assertContains(response, 'Validation Preview')
        token = self.client.session['bulk_teacher_pending']['token']
        self.client.post(reverse('bulk_upload_teachers'), {'action': 'confirm', 'upload_token': token})
        self.assertEqual(Teacher.objects.count(), 1)
        self.client.post(reverse('bulk_upload_teachers'), {'excel_file': self.upload()})
        self.client.post(reverse('bulk_upload_teachers'), {'action': 'cancel'})
        self.assertNotIn('bulk_teacher_pending', self.client.session)

    def test_template_is_importable_and_routes_are_protected(self):
        response = self.client.get(reverse('bulk_upload_teachers_template'))
        workbook = load_workbook(BytesIO(response.content))
        result = import_teachers_from_worksheet(workbook['Teachers'], preview=True)
        self.assertEqual(result.imported, 1, result.errors)
        anonymous = Client()
        for suffix in ('', '_progress', '_template', '_activity', '_credentials'):
            response = anonymous.get(reverse('bulk_upload_teachers' + suffix))
            self.assertEqual(response.status_code, 302)
            self.assertIn(response.url.split('?')[0], {reverse('login_admin'), reverse('login')})
        ordinary = User.objects.create_user('ordinary')
        self.client.force_login(ordinary)
        self.assertEqual(self.client.get(reverse('bulk_upload_teachers')).status_code, 403)

    def test_csrf_and_invalid_files(self):
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.admin)
        client.get(reverse('bulk_upload_teachers'))
        response = client.post(reverse('bulk_upload_teachers'), {'excel_file': self.upload()})
        self.assertEqual(response.status_code, 403)
        response = client.post(reverse('bulk_upload_teachers'), {'excel_file': self.upload()},
                               HTTP_X_CSRFTOKEN=client.cookies['csrftoken'].value,
                               HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 202)
        response = self.client.post(reverse('bulk_upload_teachers'),
                                    {'excel_file': SimpleUploadedFile('bad.xlsx', b'not excel')},
                                    HTTP_X_REQUESTED_WITH='XMLHttpRequest')
        self.assertEqual(response.status_code, 400)
