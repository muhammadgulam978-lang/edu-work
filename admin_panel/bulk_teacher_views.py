"""Teacher bulk uploads: validate, confirm, import in batches, and report."""
from io import BytesIO
import time
import uuid

import openpyxl
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.http import FileResponse, HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_POST

from teacher_dashboard.models import Teacher
from .models import AcademicYear, Subject
from .bulk_teacher_import import import_teachers_from_worksheet, select_teacher_worksheet

PREFIX = 'bulk_teacher_'
TEMPLATE = 'admin_panel/bulk_upload_teachers.html'


def protected(view):
    return login_required(login_url='login_admin')(
        permission_required('teacher_dashboard.add_teacher', raise_exception=True)(view)
    )


def clear_pending(request):
    pending = request.session.pop(PREFIX + 'pending', {}) or {}
    path = pending.get('path', '')
    if path.startswith(f'bulk_teacher_previews/user_{request.user.pk}/'):
        default_storage.delete(path)
    for key in ('validation', 'progress'):
        request.session.pop(PREFIX + key, None)


def context(request):
    pending = request.session.get(PREFIX + 'pending', {})
    return {
        'teacher_count': Teacher.objects.count(),
        'login_count': Teacher.objects.filter(user__is_active=True).count(),
        'active_year': AcademicYear.objects.filter(is_active=True).first(),
        'recent_teachers': Teacher.objects.select_related('user').order_by('-pk')[:6],
        'preview': pending.get('preview'), 'pending_upload': pending,
        'result': request.session.get(PREFIX + 'last_result'),
        'subjects': Subject.objects.order_by('name'),
        'faculties': Teacher.FACULTY_CHOICES,
    }


def error(request, message, status=400):
    if request.headers.get('x-requested-with') == 'XMLHttpRequest':
        return JsonResponse({'error': message}, status=status)
    messages.error(request, message)
    return redirect('bulk_upload_teachers')


def run_chunk(pending, **kwargs):
    with default_storage.open(pending['path'], 'rb') as stream:
        workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            return import_teachers_from_worksheet(workbook[pending['sheet_name']], **kwargs)
        finally:
            workbook.close()


def preview_result(result):
    return {'total_rows': result.imported + result.skipped,
            'valid_rows': result.imported, 'invalid_rows': result.skipped,
            'rows': result.rows, 'errors': result.errors, 'warnings': result.warnings}


def workbook_response(workbook, filename):
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    response = HttpResponse(stream.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def finish(request, result, started):
    credentials = result.pop('credentials', [])
    if credentials:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Teacher Logins'
        sheet.append(['Name', 'Login ID', 'Password', 'Email', 'Status'])
        for row in credentials:
            sheet.append([row[key] for key in ('name', 'username', 'password', 'email', 'status')])
            for cell in sheet[sheet.max_row]:
                cell.data_type = 's'
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        result['credentials_report_path'] = default_storage.save(
            f'bulk_teacher_credentials/user_{request.user.pk}/{uuid.uuid4().hex}.xlsx',
            ContentFile(stream.getvalue()),
        )
    result['total_rows'] = result['imported'] + result['skipped']
    result['success_rate'] = round(100 * result['imported'] / result['total_rows']) if result['total_rows'] else 0
    result['processing_seconds'] = round(time.time() - started, 2)
    request.session[PREFIX + 'last_result'] = result
    clear_pending(request)


@protected
@ensure_csrf_cookie
def bulk_upload_teachers(request):
    if request.method != 'POST':
        return render(request, TEMPLATE, context(request))
    action = request.POST.get('action', 'preview')
    pending = request.session.get(PREFIX + 'pending', {})
    if action == 'cancel':
        clear_pending(request)
        return redirect('bulk_upload_teachers')
    if action == 'confirm':
        if not request.POST.get('upload_token') or request.POST['upload_token'] != pending.get('token'):
            return error(request, 'This upload preview expired. Validate the file again.', 409)
        if not pending.get('validation_complete') or pending.get('invalid_rows') or not pending.get('valid_rows'):
            return error(request, 'Only a fully validated file can be imported.')
        if request.session.get(PREFIX + 'progress'):
            return error(request, 'Resume the import using the progress controls.', 409)
        started = time.time()
        try:
            result = run_chunk(pending)
        except Exception as exc:
            return error(request, f'Import could not continue: {exc}')
        finish(request, {key: getattr(result, key) for key in ('imported', 'skipped', 'errors', 'warnings', 'credentials', 'login_ready')}, started)
        return redirect('bulk_upload_teachers')
    if action != 'preview':
        return error(request, 'Unsupported upload action.')
    upload = request.FILES.get('excel_file')
    if not upload or not upload.name.lower().endswith('.xlsx'):
        return error(request, 'Please select an .xlsx Excel file.')
    if upload.size > getattr(settings, 'BULK_TEACHER_UPLOAD_MAX_BYTES', 10 * 1024 * 1024):
        return error(request, 'The selected file exceeds the 10 MB limit.')
    clear_pending(request)
    token = uuid.uuid4().hex
    path = default_storage.save(f'bulk_teacher_previews/user_{request.user.pk}/{token}.xlsx', upload)
    try:
        with default_storage.open(path, 'rb') as stream:
            workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
            try:
                sheet = select_teacher_worksheet(workbook)
                pending = {'token': token, 'path': path, 'name': upload.name,
                           'sheet_name': sheet.title, 'total_work_units': max(0, (sheet.max_row or 1) - 1),
                           'validation_complete': False}
            finally:
                workbook.close()
    except Exception as exc:
        default_storage.delete(path)
        return error(request, f'Could not validate this Excel file: {exc}')
    request.session[PREFIX + 'pending'] = pending
    if request.headers.get('x-requested-with') == 'XMLHttpRequest':
        return JsonResponse({'validation_pending': True, 'token': token, 'cursor': 2}, status=202)
    result = run_chunk(pending, preview=True)
    pending.update(preview_result(result))
    pending['preview'] = preview_result(result)
    pending['validation_complete'] = True
    request.session[PREFIX + 'pending'] = pending
    return render(request, TEMPLATE, context(request))


@protected
@require_POST
def bulk_upload_teachers_progress(request):
    pending = request.session.get(PREFIX + 'pending', {})
    token = request.POST.get('upload_token')
    if not token or token != pending.get('token'):
        return JsonResponse({'error': 'This upload preview expired. Validate the file again.'}, status=409)
    operation = request.POST.get('operation')
    if operation not in ('validate_start', 'validate_process', 'start', 'process'):
        return JsonResponse({'error': 'Unsupported progress operation.'}, status=400)
    validating = operation.startswith('validate_')
    if not validating and (not pending.get('validation_complete') or pending.get('invalid_rows') or not pending.get('valid_rows')):
        return JsonResponse({'error': 'Only a fully validated file can be imported.'}, status=400)
    key = PREFIX + ('validation' if validating else 'progress')
    state = request.session.get(key)
    if not state:
        if operation not in ('start', 'validate_start'):
            return JsonResponse({'error': 'Start the workflow before processing rows.'}, status=409)
        state = {'cursor': 2, 'scanned': 0, 'started': time.time(),
                 'seen': {'usernames': [], 'emails': []},
                 'result': {'imported': 0, 'skipped': 0, 'login_ready': 0,
                            'errors': [], 'warnings': [], 'rows': [], 'credentials': []}}
    done = False
    if operation in ('process', 'validate_process'):
        try:
            cursor = int(request.POST.get('cursor', 0))
        except (TypeError, ValueError):
            return JsonResponse({'error': 'Invalid progress cursor.'}, status=400)
        if cursor != state['cursor']:
            return JsonResponse({'error': 'A stale request was rejected.'}, status=409)
        try:
            chunk = run_chunk(pending, preview=validating, start_row=cursor,
                              max_rows=max(1, getattr(settings, 'BULK_TEACHER_CHUNK_ROWS', 25)),
                              seen=state['seen'] if validating else None)
        except Exception as exc:
            return JsonResponse({'error': f'Processing could not continue: {exc}'}, status=400)
        for field in ('imported', 'skipped', 'login_ready'):
            state['result'][field] += getattr(chunk, field)
        for field in ('errors', 'warnings', 'rows', 'credentials'):
            state['result'][field].extend(getattr(chunk, field))
        state['seen'] = chunk.seen if validating else state['seen']
        state['cursor'] = chunk.last_row_number + 1
        state['scanned'] += chunk.scanned_rows
        done = chunk.complete
    total = pending['total_work_units']
    elapsed = max(0.01, time.time() - state['started'])
    scanned = min(total, state['scanned'])
    request.session[key] = state
    if done:
        if validating:
            result = state['result']
            preview = {'total_rows': result['imported'] + result['skipped'], 'valid_rows': result['imported'],
                       'invalid_rows': result['skipped'], 'rows': result['rows'],
                       'errors': result['errors'], 'warnings': result['warnings']}
            pending.update(preview)
            pending.update({'preview': preview, 'validation_complete': True})
            request.session[PREFIX + 'pending'] = pending
            request.session.pop(key, None)
        else:
            state['result'].pop('rows', None)
            finish(request, state['result'], state['started'])
    return JsonResponse({
        'done': done, 'cursor': state['cursor'],
        'percent': 100 if done else min(99, round(100 * scanned / total)) if total else 0,
        'processed_rows': scanned, 'total_rows': total, 'elapsed_seconds': round(elapsed, 1),
        'eta_seconds': 0 if done else round(elapsed / scanned * (total - scanned), 1) if scanned else None,
        'redirect_url': reverse('bulk_upload_teachers'),
    })


@protected
def bulk_upload_teachers_activity(request):
    teachers = Teacher.objects.select_related('user').order_by('-pk')[:6]
    return JsonResponse({
        'teacher_count': Teacher.objects.count(),
        'login_count': Teacher.objects.filter(user__is_active=True).count(),
        'updated_at': timezone.localtime().strftime('%H:%M:%S'),
        'teachers': [{'name': teacher.name, 'email': teacher.email,
                      'department': teacher.department or 'Not assigned',
                      'url': reverse('teacher_profile', args=[teacher.pk])} for teacher in teachers],
    })


@protected
def bulk_upload_teachers_template(request):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Teachers'
    sheet.append(['Name', 'Email', 'Phone Number', 'Gender', 'Date Of Birth', 'Qualification',
                  'Experience', 'Address', 'Faculty Group', 'Department', 'Subject',
                  'Joining Date', 'Status', 'Login_Id', 'Password'])
    sheet.append(['Example Teacher', 'teacher@example.com', '03001234567', 'Female', '1990-05-06',
                  'Masters', 5, 'School Road', 'Senior Section', '', '',
                  timezone.localdate().isoformat(), 'Active', 'teacher.example', 'ChangeMe@12345'])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 24
    instructions = workbook.create_sheet('Instructions')
    for line in ['Replace the example row with your teachers. Name and Email are required. Blank Login_Id and Password values are generated.',
                 'Emails and login IDs must be unique. Passwords must satisfy the school password policy.',
                 'Separate multiple subjects or faculty groups with commas. Subjects must already exist.',
                 'Dates: YYYY-MM-DD, DD-Mon-YYYY or DD/MM/YYYY. Status: Active, Inactive or On Leave.',
                 'Id and Image columns are accepted for compatibility but are not stored.',
                 'Faculty groups: ' + ', '.join(value for value, _ in Teacher.FACULTY_CHOICES)]:
        instructions.append([line])
    instructions.column_dimensions['A'].width = 110
    return workbook_response(workbook, 'EduPilot_Teacher_Import_Template.xlsx')


@protected
def bulk_upload_teachers_credentials(request):
    path = (request.session.get(PREFIX + 'last_result') or {}).get('credentials_report_path', '')
    if not path.startswith(f'bulk_teacher_credentials/user_{request.user.pk}/') or not default_storage.exists(path):
        return error(request, 'No credentials download is available.')
    response = FileResponse(default_storage.open(path, 'rb'), as_attachment=True,
                            filename='EduPilot_Teacher_Portal_Credentials.xlsx')
    response['Cache-Control'] = 'no-store'
    return response
