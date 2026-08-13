from datetime import date

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import Workbook

from teacher_dashboard.models import Teacher

from .bulk_teacher_import import import_teachers_from_worksheet, select_teacher_worksheet
from .models import AcademicYear, Subject


class BulkTeacherImportTests(TestCase):
    def setUp(self):
        year = AcademicYear.objects.create(year='2026-27', is_active=True)
        self.subject = Subject.objects.create(
            academic_year=year,
            name='Science',
            grading_type='percentage',
            short_code='SCI-BULK',
            sort_order=1,
        )

    def worksheet(self, headers, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        return sheet

    def test_shuffled_headers_map_all_supported_teacher_fields(self):
        sheet = self.worksheet(
            [
                'Password', 'Subjects', 'Name', 'Joining Date', 'Email', 'Status',
                'Login_Id', 'Date_Of_Birth', 'Gender', 'Qualification', 'Experience',
                'Faculty Group', 'Department', 'Phone Number', 'Address',
            ],
            [[
                'StrongPass@12345', 'Science', 'Mapped Teacher', '2026-07-14',
                'mapped.teacher@example.com', 'Inactive', 'mapped.teacher', '1990-05-06',
                'Female', 'MS', 8, 'Senior Section', 'Science Department',
                '03001234567', 'School Road',
            ]],
        )

        result = import_teachers_from_worksheet(sheet)

        self.assertEqual(result.imported, 1, result.errors)
        teacher = Teacher.objects.get(email='mapped.teacher@example.com')
        self.assertEqual(teacher.name, 'Mapped Teacher')
        self.assertEqual(teacher.date_of_birth, date(1990, 5, 6))
        self.assertEqual(teacher.joining_date, date(2026, 7, 14))
        self.assertEqual(teacher.experience, 8)
        self.assertEqual(teacher.status, 'inactive')
        self.assertTrue(teacher.subjects.filter(pk=self.subject.pk).exists())
        self.assertFalse(teacher.user.is_active)

    def test_invalid_teacher_date_is_reported_without_creating_user(self):
        sheet = self.worksheet(
            ['Name', 'Email', 'Login_Id', 'Password', 'Date Of Birth'],
            [[
                'Invalid Date', 'invalid.date@example.com', 'invalid.date',
                'StrongPass@12345', 'not-a-date',
            ]],
        )

        result = import_teachers_from_worksheet(sheet)

        self.assertEqual(result.imported, 0)
        self.assertEqual(result.skipped, 1)
        self.assertIn('not a supported date', result.errors[0])
        self.assertFalse(User.objects.filter(username='invalid.date').exists())

    def test_unknown_subject_is_reported_instead_of_silently_dropped(self):
        sheet = self.worksheet(
            ['Name', 'Email', 'Login_Id', 'Password', 'Subject'],
            [['Subject Error', 'subject.error@example.com', 'subject.error', 'StrongPass@12345', 'Unknown']],
        )

        result = import_teachers_from_worksheet(sheet)

        self.assertEqual(result.imported, 0)
        self.assertEqual(result.skipped, 1)
        self.assertIn('Subject is not configured', result.errors[0])
        self.assertFalse(User.objects.filter(username='subject.error').exists())

    def test_selects_teacher_data_sheet(self):
        workbook = Workbook()
        workbook.active.append(['Instructions only'])
        data = workbook.create_sheet('Teachers')
        data.append(['Name', 'Email', 'Login_Id', 'Password'])
        data.append(['Sheet Teacher', 'sheet.teacher@example.com', 'sheet.teacher', 'StrongPass@12345'])

        selected = select_teacher_worksheet(workbook)
        result = import_teachers_from_worksheet(selected)

        self.assertEqual(selected.title, 'Teachers')
        self.assertEqual(result.imported, 1, result.errors)
