from datetime import date
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from edupilot_core.models import EmailOutbox, FeePlan, NotificationQueue, StudentFeeAssignment
from parent_dashboard.models import Parent, StudentGuardian
from student_profile.models import Student

from .bulk_student_import import (
    import_students_from_worksheet,
    preview_students_from_worksheet,
    select_student_worksheet,
)
from .bulk_credentials import get_bulk_student_password
from .models import AcademicYear, Admission, BulkStudentCredential, Class


class BulkStudentImportTests(TestCase):
    def setUp(self):
        self.year = AcademicYear.objects.create(year='2026-27', is_active=True)
        self.school_class = Class.objects.create(class_name='Grade 5')
        self.fee_plan = FeePlan.objects.create(
            name='Grade 5 Plan', class_name='Grade 5', session='2026-27'
        )

    def worksheet(self, headers, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        return sheet

    def test_optional_columns_generate_safe_defaults_and_fee_assignment(self):
        sheet = self.worksheet(
            ['S.No', 'Class', 'Fee_Plan'],
            [[1, 'Grade 5', 'Grade 5 Plan']],
        )

        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.imported, 1, result.errors)
        self.assertEqual(result.enrolled, 1)
        self.assertEqual(result.fee_ready, 1)
        student = Student.objects.get()
        self.assertEqual(student.name, 'N/A')
        self.assertTrue(student.student_id.startswith('STU-'))
        self.assertTrue(student.user.has_usable_password())
        self.assertEqual(result.student_logins_ready, 1)
        self.assertEqual(len(result.credentials), 1)
        self.assertEqual(result.parent_data_missing, 1)
        self.assertEqual(Admission.objects.get().admission_status, 'approved')
        self.assertEqual(
            StudentFeeAssignment.objects.get(canonical_student=student).fee_plan,
            self.fee_plan,
        )

    def test_duplicate_row_is_skipped_without_rolling_back_valid_rows(self):
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Email', 'Login_Id'],
            [
                ['STU-100', 'First Student', 'first@example.com', 'first.student'],
                ['STU-100', 'Duplicate Student', 'second@example.com', 'second.student'],
                ['STU-101', 'Third Student', 'third@example.com', 'third.student'],
            ],
        )

        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.imported, 2)
        self.assertEqual(result.skipped, 1)
        self.assertEqual(Student.objects.count(), 2)
        self.assertIn("Student ID 'STU-100' already exists", result.errors[0])

    def test_parent_account_is_created_linked_and_credentials_are_queued(self):
        sheet = self.worksheet(
            [
                'Student_Id', 'Student Name', 'Email', 'Class',
                'Parent_Name', 'Parent_Email', 'Parent_Phone', 'Parent_Relationship',
            ],
            [[
                'STU-200', 'Sara Ahmed', 'sara.ahmed@example.com', 'Grade 5',
                'Ahmed Raza', 'ahmed.raza@example.com', '03001112222', 'Father',
            ]],
        )

        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.imported, 1, result.errors)
        student = Student.objects.get(student_id='STU-200')
        parent = Parent.objects.get(email='ahmed.raza@example.com')
        self.assertEqual(result.parents_linked, 1)
        self.assertEqual(result.parent_accounts_created, 1)
        self.assertEqual(result.student_accounts_created, 1)
        self.assertEqual(result.student_logins_ready, 1)
        self.assertEqual(result.parent_logins_ready, 1)
        self.assertEqual(result.credential_emails_queued, 2)
        self.assertEqual(len(result.credential_email_ids), 2)
        self.assertEqual(result.messages_queued, 1)
        self.assertTrue(parent.students.filter(pk=student.pk).exists())
        self.assertTrue(parent.user.has_usable_password())
        self.assertTrue(student.user.has_usable_password())
        self.assertTrue(StudentGuardian.objects.filter(
            parent=parent, student=student, is_primary=True
        ).exists())
        self.assertEqual(EmailOutbox.objects.filter(status='PENDING').count(), 2)
        self.assertEqual(len(result.credentials), 2)

    def test_local_email_still_gets_usable_generated_portal_password(self):
        sheet = self.worksheet(
            ['Student Name', 'Email'],
            [['Local Account', 'local.account@students.edupilot.local']],
        )

        result = import_students_from_worksheet(sheet, self.year)

        student = Student.objects.get(name='Local Account')
        self.assertTrue(student.user.has_usable_password())
        self.assertEqual(result.student_logins_ready, 1)
        self.assertEqual(result.password_reset_required, 0)
        self.assertEqual(result.credentials[0]['login_id'], student.user.username)

    def test_each_bulk_student_gets_a_distinct_securely_stored_password(self):
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Login_Id', 'Password'],
            [
                ['STU-PASS-1', 'Similar Student', 'similar.student.1', 'SharedSheetPassword'],
                ['STU-PASS-2', 'Similar Student', 'similar.student.2', 'SharedSheetPassword'],
            ],
        )

        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.imported, 2, result.errors)
        passwords = [item['password'] for item in result.credentials]
        self.assertEqual(len(passwords), len(set(passwords)))
        self.assertEqual(BulkStudentCredential.objects.count(), 2)
        for item in result.credentials:
            student = Student.objects.get(student_id=item['student_id'])
            self.assertTrue(student.user.check_password(item['password']))
            self.assertEqual(get_bulk_student_password(student), item['password'])

    def test_parent_data_is_not_fabricated_when_workbook_does_not_supply_it(self):
        sheet = self.worksheet(['Student Name'], [['Student Without Parent']])

        preview = preview_students_from_worksheet(sheet, self.year)
        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(preview.parent_data_missing, 1)
        self.assertEqual(result.parent_data_missing, 1)
        self.assertEqual(Parent.objects.count(), 0)

    def test_student_phone_creates_trackable_message_queue_record(self):
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Email', 'Contact No'],
            [['STU-202', 'Message Student', 'message.student@example.com', '03001234567']],
        )

        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.imported, 1, result.errors)
        self.assertEqual(result.messages_queued, 1)
        self.assertEqual(len(result.message_ids), 1)
        queued = NotificationQueue.objects.get(pk=result.message_ids[0])
        self.assertEqual(queued.status, 'PENDING')
        self.assertEqual(queued.canonical_student.student_id, 'STU-202')

    def test_existing_parent_is_reused_without_password_change(self):
        from django.contrib.auth.models import User

        user = User.objects.create_user(
            username='existing.parent', email='parent@example.com', password='KeepMe@12345'
        )
        parent = Parent.objects.create(
            user=user, full_name='Existing Parent', email='parent@example.com'
        )
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Parent_Email', 'Parent_Password'],
            [['STU-201', 'Second Child', 'parent@example.com', 'DoNotUse@12345']],
        )

        result = import_students_from_worksheet(sheet, self.year)

        user.refresh_from_db()
        self.assertEqual(Parent.objects.count(), 1)
        self.assertEqual(result.parent_accounts_created, 0)
        self.assertTrue(user.check_password('KeepMe@12345'))
        self.assertTrue(parent.students.filter(student_id='STU-201').exists())

    def test_preview_reports_student_parent_and_fee_readiness_without_writes(self):
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Class', 'Parent_Name', 'Parent_Email'],
            [['STU-300', 'Preview Student', 'Grade 5', 'Preview Parent', 'preview.parent@example.com']],
        )

        result = preview_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.total_rows, 1)
        self.assertEqual(result.valid_rows, 1)
        self.assertEqual(result.parent_rows, 1)
        self.assertEqual(result.new_parents, 1)
        self.assertEqual(result.fee_ready_rows, 1)
        self.assertEqual(Student.objects.count(), 0)
        self.assertEqual(Parent.objects.count(), 0)

    def test_preview_blocks_duplicates_before_import(self):
        Student.objects.create(
            student_id='STU-301', name='Existing', father_name='N/A', mother_name='N/A',
            roll_no='301', date_of_birth=date.today(),
            email='existing.student@example.com',
        )
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Email'],
            [['STU-301', 'Duplicate', 'duplicate@example.com']],
        )

        result = preview_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.invalid_rows, 1)
        self.assertEqual(result.duplicate_rows, 1)
        self.assertIn("Student ID 'STU-301' already exists", result.errors[0])

    def test_parent_dataset_child_columns_create_named_student(self):
        sheet = self.worksheet(
            [
                'S.No', 'Parent ID', 'Father Name', 'Father Phone', 'Father Email',
                'Guardian Name', 'Guardian Relationship', 'Guardian Phone',
                'Home Address', 'Child Name', 'Child Admission Number',
                'Registration Date', 'Is Active',
            ],
            [[
                1, 'PAR5002', 'Ahmed Ahmed', '03238719076',
                'ahmed.ahmed1@edupilot.test', 'Ismail Ahmed', 'Mother',
                '03167303868', 'House 982, Street 43', 'Hassan Ahmed',
                'STU5002', '2025-10-16', 'Yes',
            ]],
        )

        preview = preview_students_from_worksheet(sheet, self.year)
        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(preview.rows[0]['name'], 'Hassan Ahmed')
        self.assertEqual(preview.rows[0]['student_id'], 'STU5002')
        self.assertEqual(result.imported, 1, result.errors)
        student = Student.objects.get(student_id='STU5002')
        self.assertEqual(student.name, 'Hassan Ahmed')
        self.assertEqual(student.address, 'House 982, Street 43')
        self.assertEqual(Admission.objects.get(student_id='STU5002').name, 'Hassan Ahmed')

    def test_selects_data_sheet_instead_of_active_instructions_sheet(self):
        workbook = Workbook()
        instructions = workbook.active
        instructions.title = 'Instructions'
        instructions.append(['Read this file before importing'])
        data = workbook.create_sheet('Student Data')
        data.append(['Student_Id', 'Student Name', 'Email'])
        data.append(['STU-SHEET-1', 'Correct Sheet Student', 'sheet.student@example.com'])

        selected = select_student_worksheet(workbook)
        result = import_students_from_worksheet(selected, self.year)

        self.assertEqual(selected.title, 'Student Data')
        self.assertEqual(result.imported, 1, result.errors)
        self.assertEqual(Student.objects.get().name, 'Correct Sheet Student')

    def test_numeric_identifiers_do_not_gain_decimal_suffix(self):
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Roll_No'],
            [[5001.0, 'Numeric Student', 17.0]],
        )

        result = import_students_from_worksheet(sheet, self.year)

        self.assertEqual(result.imported, 1, result.errors)
        student = Student.objects.get()
        self.assertEqual(student.student_id, '5001')
        self.assertEqual(student.roll_no, '17')

    def test_invalid_row_does_not_reserve_identifiers_from_later_valid_row(self):
        sheet = self.worksheet(
            ['Student_Id', 'Student Name', 'Email', 'Login_Id', 'Password'],
            [
                ['STU-REUSE-1', 'Invalid First', 'reuse@example.com', 'reuse.login', 'short'],
                ['STU-REUSE-1', 'Valid Second', 'reuse@example.com', 'reuse.login', 'StrongPass@12345'],
            ],
        )

        preview = preview_students_from_worksheet(sheet, self.year)

        self.assertEqual(preview.invalid_rows, 1)
        self.assertEqual(preview.valid_rows, 1)

    def test_duplicate_header_is_rejected_explicitly(self):
        sheet = self.worksheet(
            ['Student Name', 'Student Name', 'Email'],
            [['First Value', 'Second Value', 'duplicate.header@example.com']],
        )

        with self.assertRaisesRegex(ValueError, 'Duplicate Excel headers'):
            preview_students_from_worksheet(sheet, self.year)

    def test_teacher_workbook_is_rejected_by_student_import(self):
        sheet = self.worksheet(
            [
                'S.No', 'Teacher Name', 'Teacher ID', 'Department', 'Subject',
                'Qualification', 'Experience (Years)', 'Phone Number', 'Email',
                'Joining Date', 'Is Active',
            ],
            [[
                1, 'Wrong Portal Teacher', 'T1001', 'Science', 'Science',
                'Masters', 5, '03001234567', 'teacher@example.com',
                '2026-01-01', 'Yes',
            ]],
        )

        with self.assertRaisesRegex(ValueError, 'teacher-format workbook'):
            preview_students_from_worksheet(sheet, self.year)

    def test_student_sheet_is_selected_when_workbook_also_has_teacher_sheet(self):
        workbook = Workbook()
        teacher_sheet = workbook.active
        teacher_sheet.title = 'Teachers'
        teacher_sheet.append(['Teacher Name', 'Teacher ID', 'Department', 'Email'])
        teacher_sheet.append(['Teacher One', 'T1', 'Science', 'teacher@example.com'])
        student_sheet = workbook.create_sheet('Students')
        student_sheet.append(['Student_Id', 'Student Name', 'Email'])
        student_sheet.append(['STU-MIXED-1', 'Correct Student', 'student@example.com'])

        selected = select_student_worksheet(workbook)

        self.assertEqual(selected.title, 'Students')


class BulkStudentLoginActivationTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(
            username='bulk-admin', email='bulk-admin@example.com', password='AdminPass@123'
        )
        self.client.force_login(self.admin)
        self.year = AcademicYear.objects.create(year='2027-28', is_active=True)
        self.student_ids = []
        for index in range(10):
            user = User.objects.create_user(
                username=f'student_stu_batch_{index}',
                email=f'batch{index}@students.edupilot.local',
            )
            user.set_unusable_password()
            user.save(update_fields=['password'])
            student = Student.objects.create(
                academic_year=self.year,
                user=user,
                student_id=f'STU-BATCH-{index}',
                name=f'Batch Student {index}',
                father_name='N/A',
                mother_name='N/A',
                roll_no=str(index + 1),
                gender='Male',
                date_of_birth=date(2012, 1, 1),
                email=f'batch{index}@students.edupilot.local',
            )
            self.student_ids.append(student.student_id)
        session = self.client.session
        session['bulk_student_last_result'] = {
            'student_ids': self.student_ids,
            'student_accounts_created': 10,
            'student_logins_ready': 0,
            'parent_logins_ready': 0,
            'password_reset_required': 10,
        }
        session.save()

    @patch('admin_panel.views._save_bulk_student_credentials', return_value='bulk/credentials.xlsx')
    def test_activation_runs_in_short_batches_and_finishes_report(self, save_credentials):
        url = reverse('bulk_upload_students_activate_logins')
        start = self.client.post(
            url, {'operation': 'start'}, HTTP_X_REQUESTED_WITH='XMLHttpRequest'
        )

        self.assertEqual(start.status_code, 200)
        self.assertEqual(start.json()['total_rows'], 10)

        first = self.client.post(
            url, {'operation': 'process', 'cursor': 0},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertFalse(first.json()['done'])
        self.assertEqual(first.json()['processed_rows'], 8)

        second = self.client.post(
            url, {'operation': 'process', 'cursor': 8},
            HTTP_X_REQUESTED_WITH='XMLHttpRequest',
        )
        self.assertTrue(second.json()['done'])
        self.assertEqual(second.json()['processed_rows'], 10)
        self.assertTrue(all(student.user.has_usable_password() for student in Student.objects.select_related('user')))
        report = self.client.session['bulk_student_last_result']
        self.assertEqual(report['student_logins_ready'], 10)
        self.assertEqual(report['password_reset_required'], 0)
        self.assertEqual(report['credentials_report_path'], 'bulk/credentials.xlsx')
        self.assertEqual(len(save_credentials.call_args.args[1]), 10)
